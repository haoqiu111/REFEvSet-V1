"""Small drawing layer on python-pptx for framework figures in the style of a fusion / detection pipeline diagram:
perspective image stacks as illustrations, 3-D slab blocks for modules, tall bars for losses, circular operator nodes,
thick translucent arrows (green: forward path, blue: gradient / back-propagation path, dashed: reference and control
path), dashed region frames and italic serif labels. Every element is a native, editable PowerPoint shape.
export(pptx) renders the slide with PowerPoint itself into a vector PDF and a PNG, so the manuscript figure is the
editable file."""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
GREEN, BLUE, ORANGE, GREY, DARK = 'A9D18E', '9DC3E6', 'F4B183', '7F7F7F', '262626'
ORANGE_D, BLUE_D, GREEN_D = 'C55A11', '2E75B6', '548235'
FONT = 'Times New Roman'


def rgb(h): return RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:], 16))


class Canvas:
    def __init__(self, w, h):
        self.prs = Presentation(); self.prs.slide_width = Inches(w); self.prs.slide_height = Inches(h); self.W, self.H = w, h
        self.s = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    # ---- text -------------------------------------------------------------------------------------------------
    def text(self, x, y, w, h, lines, size=11, italic=True, bold=False, color=DARK, align='c', anchor='m'):
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {'m': MSO_ANCHOR.MIDDLE, 't': MSO_ANCHOR.TOP, 'b': MSO_ANCHOR.BOTTOM}[anchor]
        for i, ln in enumerate([lines] if isinstance(lines, str) else lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = {'c': PP_ALIGN.CENTER, 'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT}[align]
            self._runs(p, ln, size, italic, bold, color)
        return tb

    def _runs(self, p, ln, size, italic, bold, color):
        """'L_{CE}' style subscripts: text inside _{...} is written as a subscript run"""
        import re
        for tok in re.split(r'(_\{[^}]*\}|\^\{[^}]*\})', ln):
            if not tok: continue
            r = p.add_run(); sub = tok.startswith('_{'); sup = tok.startswith('^{'); r.text = tok[2:-1] if (sub or sup) else tok
            r.font.size = Pt(size); r.font.italic = italic; r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color)
            if sub: r.font._element.set('baseline', '-25000')
            if sup: r.font._element.set('baseline', '30000')

    def _style(self, shp, fill, line, lw=0.75, alpha=None, dash=None):
        if fill is None: shp.fill.background()
        else:
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
            if alpha is not None:
                sf = shp.fill._xPr.find(qn('a:solidFill'))[0]; a = etree.SubElement(sf, f'{{{A}}}alpha'); a.set('val', str(int(alpha * 100000)))
        if line is None: shp.line.fill.background()
        else:
            shp.line.color.rgb = rgb(line); shp.line.width = Pt(lw)
            if dash:
                ln = shp.line._get_or_add_ln(); d = etree.SubElement(ln, f'{{{A}}}prstDash'); d.set('val', dash)
        shp.shadow.inherit = False

    def _label(self, shp, lines, size, italic=True, bold=False, color=DARK):
        tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.03); tf.margin_top = tf.margin_bottom = Inches(0.02)
        for i, ln in enumerate([lines] if isinstance(lines, str) else lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER; self._runs(p, ln, size, italic, bold, color)

    # ---- shapes -----------------------------------------------------------------------------------------------
    def frame(self, x, y, w, h, dash='dash', color=GREY, lw=0.75):
        shp = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); self._style(shp, None, color, lw, dash=dash); return shp

    def box(self, x, y, w, h, lines='', fill='FFFFFF', line=GREY, size=10, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=0.75, alpha=None, italic=True, bold=False, color=DARK, dash=None):
        shp = self.s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); self._style(shp, fill, line, lw, alpha, dash)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE: shp.adjustments[0] = 0.12
        if lines: self._label(shp, lines, size, italic, bold, color)
        return shp

    def cube(self, x, y, w, h, lines='', fill=ORANGE, line=ORANGE_D, size=10, depth=0.25, alpha=None, color=DARK, bold=False):
        """3-D block (module); depth = share of the smaller side used for the top / side faces"""
        shp = self.s.shapes.add_shape(MSO_SHAPE.CUBE, Inches(x), Inches(y), Inches(w), Inches(h)); shp.adjustments[0] = depth; self._style(shp, fill, line, 0.75, alpha)
        if lines: self._label(shp, lines, size, True, bold, color)
        return shp

    def slabs(self, x, y, w, h, n=2, fill=ORANGE, line=ORANGE_D, gap=0.16, depth=0.5):
        """n thin upright slabs behind each other (encoder / decoder look)"""
        out = []
        for i in range(n):
            out.append(self.cube(x + i * gap, y + (n - 1 - i) * gap * 0.55, w, h - (n - 1) * gap * 0.55, fill=fill, line=line, depth=depth, alpha=0.85))
        return out

    def node(self, x, y, d, sym='+', fill='FFFFFF', line=BLUE_D, size=14, color=BLUE_D):
        shp = self.s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - d / 2), Inches(y - d / 2), Inches(d), Inches(d)); self._style(shp, fill, line, 1.5)
        self._label(shp, sym, size, italic=False, bold=True, color=color); return shp

    def bar(self, x, y, w, h, lines, fill, size=10, color=DARK):
        """tall flat bar (loss)"""
        shp = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); self._style(shp, fill, None, alpha=0.9)
        self._label(shp, lines, size, True, False, color); return shp

    def picture(self, path, x, y, w, h=None, camera=None, border=None):
        pic = self.s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h) if h else None)
        sp = pic._element.spPr
        if border:
            ln = etree.SubElement(sp, f'{{{A}}}ln'); ln.set('w', str(int(Pt(0.5)))); sf = etree.SubElement(ln, f'{{{A}}}solidFill'); c = etree.SubElement(sf, f'{{{A}}}srgbClr'); c.set('val', border)
        if camera:
            sc = etree.SubElement(sp, f'{{{A}}}scene3d'); cam = etree.SubElement(sc, f'{{{A}}}camera'); cam.set('prst', camera)
            lr = etree.SubElement(sc, f'{{{A}}}lightRig'); lr.set('rig', 'flat'); lr.set('dir', 't')
        return pic

    def stack(self, paths, x, y, w, h, dx=0.13, dy=-0.07, camera='isometricOffAxis1Right'):
        """perspective stack of images (back to front)"""
        n = len(paths)
        for i, p in enumerate(paths):
            self.picture(p, x + (n - 1 - i) * dx, y + (n - 1 - i) * dy, w, h, camera=camera, border='BFBFBF')

    # ---- arrows -----------------------------------------------------------------------------------------------
    def arrow(self, x1, y1, x2, y2, color=GREEN, lw=4.5, kind='straight', dash=None, head=True, alpha=None):
        c = self.s.shapes.add_connector({'straight': MSO_CONNECTOR.STRAIGHT, 'elbow': MSO_CONNECTOR.ELBOW, 'curve': MSO_CONNECTOR.CURVE}[kind], Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color); c.line.width = Pt(lw); ln = c.line._get_or_add_ln()
        if alpha is not None:
            a = etree.SubElement(ln.find(qn('a:solidFill'))[0], f'{{{A}}}alpha'); a.set('val', str(int(alpha * 100000)))
        if dash:
            d = etree.SubElement(ln, f'{{{A}}}prstDash'); d.set('val', dash)
        if head:
            t = etree.SubElement(ln, f'{{{A}}}tailEnd'); t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
        return c

    def path(self, pts, color=GREEN, lw=4.5, dash=None, alpha=None):
        """poly-line arrow through the points (the last segment carries the head)"""
        for i in range(len(pts) - 1):
            self.arrow(*pts[i], *pts[i + 1], color=color, lw=lw, dash=dash, head=(i == len(pts) - 2), alpha=alpha)

    def save(self, path):
        self.prs.save(path); return path


def export(pptx_path, width_px=4200):
    """PowerPoint renders the slide: <name>.pdf (vector) and <name>.png"""
    import win32com.client
    pptx_path = os.path.abspath(pptx_path); base = os.path.splitext(pptx_path)[0]
    app = win32com.client.Dispatch('PowerPoint.Application'); pres = app.Presentations.Open(pptx_path, WithWindow=False)
    try:
        pres.SaveAs(base + '.pdf', 32)
        sw, sh = pres.PageSetup.SlideWidth, pres.PageSetup.SlideHeight
        pres.Slides(1).Export(base + '.png', 'PNG', width_px, int(round(width_px * sh / sw)))
    finally:
        pres.Close()
    return base + '.pdf', base + '.png'
